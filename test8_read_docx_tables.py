import docx
import pandas as pd


file=docx.Document("shezhen_test.docx")
tables = file.tables #获取文件中的表格集

print("len=",len(tables))

for idx,table in enumerate(tables):
    # print(table)
    print(f"Table {idx}:")
    if idx==0:

        name=table.cell(0,1).text
        age=table.cell(1,1).text
        sex=table.cell(0,3).text
        report_date=table.cell(1,3).text
        print(name,age,sex,report_date)
        print()
    if idx==1:
        tongue_color=table.cell(1,2).text
        tongue_tip=table.cell(2,2).text
        tongue_shape=table.cell(3,2).text
        tongue_coating_color=table.cell(4,2).text
        tongue_coating_nature=table.cell(5,2).text
        tongue_other_features=table.cell(6,1).text
        print(tongue_color,tongue_tip,tongue_shape,
              tongue_coating_color,tongue_coating_nature,
              tongue_other_features)
        print()

    if idx==2:
        tongue_color_index = table.cell(2,1).text
        tongue_color_feature = table.cell(2,2).text
        tongue_color_reference = table.cell(2,3).text
        coating_color_index = table.cell(3,1).text
        coating_color_feature = table.cell(3,2).text
        coating_color_reference = table.cell(3,3).text

        obesity_index = table.cell(4,1).text
        obesity_feature = table.cell(4,2).text
        obesity_reference = table.cell(4,3).text

        thick_thin_index = table.cell(5,1).text
        thick_thin_feature = table.cell(5,2).text
        thick_thin_reference = table.cell(5,3).text

        rotten_index = table.cell(6,1).text
        rotten_feature = table.cell(6,2).text
        rotten_reference = table.cell(6,3).text
        print(tongue_color_index,tongue_color_feature,tongue_color_reference,'\n',
              coating_color_index,coating_color_feature,coating_color_reference,'\n',
              obesity_index,obesity_feature,obesity_reference,'\n',
              thick_thin_index,thick_thin_feature,obesity_reference,'\n',
              rotten_index,thick_thin_feature,thick_thin_reference)
